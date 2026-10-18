# SPDX-License-Identifier: AGPL-3.0-or-later
"""The DOCX export of a version (X03c, ADR-0036), built with python-docx: headings, paragraphs
and their formatting (fonts, sizes, colours, highlights, strike, sub- and superscript, links,
alignment), lists, quotations, boxes, separators, images at their width (the text flowing around
an image at the left or right, as Word's square wrapping), the Sources block and the AI mention."""

from __future__ import annotations

import copy
import io
from collections.abc import Callable
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from otlet.writing.domain.document import FONTS, MENTION_IA, SOURCES_TITLE, ai_share

Node = dict[str, Any]
# The image of a media id: its bytes, width and height in pixels; None when it is gone.
MediaBytes = Callable[[str], tuple[bytes, int, int] | None]

LINK = RGBColor(0x2F, 0x5D, 0x8A)
QUIET = RGBColor(0x5C, 0x5C, 0x57)
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
COLUMN_CM = 21.0 - 2 * 2.2  # A4 less the margins
WRAP_DEFAULT = 50


def _shade(element: Any, colour: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), colour.lstrip("#").upper())
    element.append(shading)


def _runs(paragraph: Paragraph, content: list[Node]) -> None:
    for node in content:
        marks = {m["type"]: m.get("attrs") or {} for m in node.get("marks", [])}
        run = paragraph.add_run(node["text"])
        font = run.font
        font.bold = "bold" in marks or None
        font.italic = "italic" in marks or None
        font.strike = "strike" in marks or None
        font.subscript = "subscript" in marks or None
        font.superscript = "superscript" in marks or "source_ref" in marks or None
        style = marks.get("text_style", {})
        if style.get("font") in FONTS:
            font.name = FONTS[style["font"]].split(",")[0].strip("'\" ")
        if style.get("size"):
            font.size = Pt(style["size"])
        if style.get("color"):
            font.color.rgb = RGBColor.from_string(style["color"].lstrip("#").upper())
        if "highlight" in marks:
            _shade(run._r.get_or_add_rPr(), marks["highlight"]["color"])
        if "link" in marks or "source_ref" in marks:
            font.color.rgb = LINK
        if "link" in marks:
            font.underline = True
            relation = paragraph.part.relate_to(
                marks["link"]["href"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), relation)
            run._r.addprevious(link)
            link.append(run._r)


def _align(paragraph: Paragraph, node: Node) -> None:
    align = (node.get("attrs") or {}).get("align")
    if align in ALIGN:
        paragraph.alignment = ALIGN[align]


def _list(document: Any, node: Node, depth: int) -> None:
    kind = "List Number" if node["type"] == "ordered_list" else "List Bullet"
    style = kind if depth == 0 else f"{kind} {min(depth + 1, 3)}"
    for item in node["content"]:
        for child in item["content"]:
            if child["type"] == "paragraph":
                paragraph = document.add_paragraph(style=style)
                _runs(paragraph, child.get("content", []))
                _align(paragraph, child)
            else:
                _list(document, child, depth + 1)


def _anchor(run: Any, side: str) -> None:
    """Turns an inline picture into one the text flows around, at the left or the right of the
    column (Word's square wrapping)."""
    inline = run._r.xpath(".//wp:inline")[0]
    extent = inline.xpath("./wp:extent")[0]
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="114300" distL="114300"'
        ' distR="114300" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0"'
        ' layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:align>{side}</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{extent.get("cx")}" cy="{extent.get("cy")}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/>'
        "</wp:anchor>"
    )
    for name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(copy.deepcopy(inline.xpath(f"./{name}")[0]))
    inline.addprevious(anchor)
    inline.getparent().remove(inline)


def _image(document: Any, block: Node, media: MediaBytes) -> None:
    attrs = block["attrs"]
    found = media(attrs["media"])
    if found is None:
        return
    data, width, height = found
    side = attrs.get("align") if attrs.get("wrap") else None
    share = attrs.get("width") or (WRAP_DEFAULT if side else None)
    # Its width in the column; else its own (at 96 dpi), never wider than the column.
    wide = COLUMN_CM * share / 100 if share else min(COLUMN_CM, width / 96 * 2.54)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(data), width=Cm(wide), height=Cm(wide * height / width))
    if side in ("left", "right"):
        _anchor(run, side)
    else:
        _align(paragraph, block)


def _separator(document: Any) -> None:
    paragraph = document.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "B0B0A8")):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def docx(title: str, document: Node, media: MediaBytes) -> bytes:
    out = Document()
    section = out.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.2))
    section.top_margin = section.bottom_margin = Cm(2)
    normal = out.styles["Normal"].font
    normal.name, normal.size = "Georgia", Pt(12)
    out.add_paragraph(title, style="Title")
    sources = {s["id"]: s for s in document["attrs"]["sources"]}
    for block in document["content"]:
        kind = block["type"]
        if kind == "heading":
            paragraph = out.add_paragraph(style=f"Heading {block['attrs']['level'] - 1}")
            _runs(paragraph, block.get("content", []))
            _align(paragraph, block)
        elif kind == "paragraph":
            paragraph = out.add_paragraph()
            _runs(paragraph, block.get("content", []))
            _align(paragraph, block)
        elif kind in ("bullet_list", "ordered_list"):
            _list(out, block, 0)
        elif kind == "citation":
            for child in block["content"]:
                paragraph = out.add_paragraph(style="Quote")
                _runs(paragraph, child.get("content", []))
            source = sources.get(block["attrs"]["source"], {"titre": "?"})
            attribution = out.add_paragraph()
            lead = attribution.add_run("— ")
            lead.font.color.rgb = QUIET
            link = [{"type": "link", "attrs": {"href": source["url"]}}] if source.get("url") else []
            _runs(attribution, [{"type": "text", "text": source["titre"], "marks": link}])
        elif kind == "encadre":
            for child in block["content"]:
                paragraph = out.add_paragraph()
                _shade(paragraph._p.get_or_add_pPr(), "#F2F2EE")
                _runs(paragraph, child.get("content", []))
        elif kind == "separateur":
            _separator(out)
        elif kind == "image":
            _image(out, block, media)
    if ai_share(document) > 0:
        mention = out.add_paragraph().add_run(MENTION_IA)
        mention.font.italic = True
        mention.font.color.rgb = QUIET
    if sources:
        out.add_paragraph(SOURCES_TITLE, style="Heading 1")
        for n, source in enumerate(sources.values(), 1):
            paragraph = out.add_paragraph()
            paragraph.add_run(f"{n}. ")
            link = [{"type": "link", "attrs": {"href": source["url"]}}] if source.get("url") else []
            _runs(paragraph, [{"type": "text", "text": source["titre"], "marks": link}])
            details = ", ".join(v for v in (source.get("editeur"), source.get("date")) if v)
            if details:
                paragraph.add_run(f", {details}")
    stream = io.BytesIO()
    out.save(stream)
    return stream.getvalue()
