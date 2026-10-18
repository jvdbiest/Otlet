# SPDX-License-Identifier: AGPL-3.0-or-later
"""Exports of a version (X03c): every format carries the text, its formatting where the format
has it, the images, the Sources block and the AI mention."""

import io
import json
import zipfile
from datetime import UTC, datetime
from typing import Any
from uuid import UUID, uuid4

import pytest
from docx import Document
from PIL import Image as PilImage

from otlet.workspaces.domain.model import Role, WorkspaceContext
from otlet.writing.adapters.export_docx import docx
from otlet.writing.application.exports import Exports
from otlet.writing.domain.article import Version
from otlet.writing.domain.export import filename

IMAGE = "0b0e6f5e-3c4e-4c33-9a54-6f1f3c1b2a10"
SOURCES = [
    {
        "id": "S1",
        "titre": "Le journal",
        "url": "https://journal.example.org/1",
        "editeur": "Journal",
    }
]


def png() -> bytes:
    out = io.BytesIO()
    PilImage.new("RGB", (400, 200), (20, 120, 200)).save(out, "PNG")
    return out.getvalue()


def text(value: str, *marks: dict[str, Any]) -> dict[str, Any]:
    node: dict[str, Any] = {"type": "text", "text": value}
    if marks:
        node["marks"] = list(marks)
    return node


DOCUMENT: dict[str, Any] = {
    "type": "doc",
    "attrs": {"version_schema": 1, "langue": "fr", "sources": SOURCES},
    "content": [
        {
            "type": "paragraph",
            "attrs": {"id": "b1", "origine": "ia", "align": "center"},
            "content": [
                text("Un chapeau "),
                text("gras", {"type": "bold"}),
                text(" et "),
                text("rouge", {"type": "text_style", "attrs": {"color": "#c00000", "size": 14}}),
                text(" [1]", {"type": "source_ref", "attrs": {"source": "S1"}}),
            ],
        },
        {
            "type": "heading",
            "attrs": {"id": "b2", "origine": "ia", "level": 2},
            "content": [text("Intertitre")],
        },
        {
            "type": "bullet_list",
            "attrs": {"id": "b3", "origine": "humain"},
            "content": [
                {
                    "type": "list_item",
                    "content": [{"type": "paragraph", "content": [text("un point")]}],
                }
            ],
        },
        {
            "type": "citation",
            "attrs": {"id": "b4", "origine": "humain", "source": "S1"},
            "content": [{"type": "paragraph", "content": [text("Une phrase citée")]}],
        },
        {
            "type": "image",
            "attrs": {
                "id": "b5",
                "origine": "humain",
                "media": IMAGE,
                "alt": "Une carte",
                "align": "left",
                "wrap": True,
                "width": 40,
            },
        },
        {
            "type": "paragraph",
            "attrs": {"id": "b6", "origine": "humain"},
            "content": [
                text("Un lien", {"type": "link", "attrs": {"href": "https://exemple.org"}})
            ],
        },
    ],
}


class Store:
    def version(self, workspace_id: UUID, article_id: UUID, number: int) -> Version | None:
        if number != 1:
            return None
        return Version(
            1,
            "Max Verstappen : un pilote",
            DOCUMENT,
            "humain",
            None,
            None,
            None,
            (),
            (),
            False,
            datetime.now(UTC),
        )


def exports(pdf: Any = lambda html: b"%PDF" + html.encode()[:10]) -> Exports:
    data = png()
    return Exports(
        Store(),
        lambda context, media: (data, "image/png", 400, 200) if media == IMAGE else None,
        docx,
        pdf,
    )


CONTEXT = WorkspaceContext(user_id=uuid4(), workspace_id=uuid4(), role=Role.OWNER)


def test_a_file_name_comes_from_the_title() -> None:
    assert (
        filename("Max Verstappen : un pilote d'exception", "pdf")
        == "max-verstappen-un-pilote-d-exception.pdf"
    )
    assert filename("???", "md") == "article.md"


def test_html_is_a_whole_page_with_its_images_inside() -> None:
    exported = exports().export(CONTEXT, uuid4(), 1, "html")
    assert exported is not None and exported.filename == "max-verstappen-un-pilote.html"
    page = exported.data.decode()
    assert page.startswith("<!doctype html>") and "<h1>Max Verstappen : un pilote</h1>" in page
    assert 'src="data:image/png;base64,' in page and "float: left" in page
    assert "Cet article a été rédigé avec l’aide" in page and "otlet-sources-title" in page
    assert exports().export(CONTEXT, uuid4(), 2, "html") is None


def test_markdown_with_images_is_a_zip_text_alone_has_none() -> None:
    exported = exports().export(CONTEXT, uuid4(), 1, "md")
    assert exported is not None and exported.mime == "application/zip"
    with zipfile.ZipFile(io.BytesIO(exported.data)) as bundle:
        names = bundle.namelist()
        written = bundle.read("max-verstappen-un-pilote.md").decode()
    assert f"images/{IMAGE}.png" in names
    assert "**gras**" in written and "## Intertitre" in written and "- un point" in written
    assert f"![Une carte](images/{IMAGE}.png)" in written and "> Une phrase citée" in written
    assert "1. [Le journal](https://journal.example.org/1), Journal" in written
    plain = exports().export(CONTEXT, uuid4(), 1, "txt")
    assert plain is not None
    body = plain.data.decode()
    assert "Un chapeau gras et rouge [1]" in body and "[Image : Une carte]" in body
    assert "« Une phrase citée »" in body and "https://journal.example.org/1" in body
    as_json = exports().export(CONTEXT, uuid4(), 1, "json")
    assert as_json is not None and json.loads(as_json.data)["document"] == DOCUMENT


def test_docx_keeps_formatting_images_and_sources() -> None:
    exported = exports().export(CONTEXT, uuid4(), 1, "docx")
    assert exported is not None
    read = Document(io.BytesIO(exported.data))
    texts = [p.text for p in read.paragraphs]
    assert texts[0] == "Max Verstappen : un pilote"
    lead = read.paragraphs[1]
    assert lead.alignment == 1  # centred
    bold = [r for r in lead.runs if r.text == "gras"][0]
    red = [r for r in lead.runs if r.text == "rouge"][0]
    assert bold.bold and str(red.font.color.rgb) == "C00000" and red.font.size is not None
    assert red.font.size.pt == 14
    assert "Intertitre" in texts and "un point" in texts and "Une phrase citée" in texts
    xml = read.element.xml
    assert "wp:anchor" in xml and "wrapSquare" in xml  # the text flows around the image
    assert "Sources" in texts and any(t.startswith("1. Le journal") for t in texts)
    assert read.part.rels and any(r.reltype.endswith("/hyperlink") for r in read.part.rels.values())


def test_pdf_is_printed_from_the_html_page() -> None:
    seen: list[str] = []

    def printer(html: str) -> bytes:
        seen.append(html)
        return b"%PDF-1.7"

    exported = exports(printer).export(CONTEXT, uuid4(), 1, "pdf")
    assert exported is not None and exported.data == b"%PDF-1.7" and "@page" in seen[0]


def test_weasyprint_prints_a_real_pdf() -> None:
    try:
        from otlet.writing.adapters.export_pdf import pdf

        data = pdf("<html><body><p>Bonjour</p></body></html>")
    except OSError as error:  # no Pango on this machine (Windows CI, slim test images)
        pytest.skip(f"WeasyPrint cannot load Pango here: {error}")
    assert data.startswith(b"%PDF")
